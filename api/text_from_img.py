import os
import platform
import pytesseract
import fitz  
from PIL import Image
from docx import Document
import openpyxl
from io import BytesIO

if platform.system() == "Windows":
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

class ImageToText:
    def __init__(self):
        self.text = ""

    def extract_txt_from_file(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()

        if ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.webp']:
            self._extract_from_image(file_path)
  
        elif ext == '.pdf':
            self._extract_from_pdf(file_path)
            
        
        elif ext == '.docx':
            self._extract_from_docx(file_path)
    
        elif ext == '.doc':
            self._extract_from_doc_via_pdf(file_path)
            
        # 5. Matn va jadval fayllari (.txt, .csv)
        elif ext in ['.txt', '.csv']:
            self._extract_from_text_or_csv(file_path)
            
        # 6. Excel (.xlsx) fayllar
        elif ext == '.xlsx':
            self._extract_from_excel(file_path)
            
        else:
            raise ValueError(f"Qo'llab-quvvatlanmaydigan fayl formati: {ext}")
        
        return self.text

    def _extract_from_image(self, image_path):
        image = Image.open(image_path)
        self.text = pytesseract.image_to_string(image, lang="eng")

    def _extract_from_pdf(self, pdf_path):
        """PDF sahifalarinima-sahifa rasmga o'girib, Tesseract'dan o'tkazadi"""
        extracted_text = []
        doc = fitz.open(pdf_path)
        
        for page_index in range(len(doc)):
            page = doc[page_index]
            pix = page.get_pixmap(dpi=150)
            img_data = pix.tobytes("png")
            image = Image.open(BytesIO(img_data))
            
            page_text = pytesseract.image_to_string(image, lang="eng")
            extracted_text.append(page_text)
            
        self.text = "\n".join(extracted_text)
        doc.close()

    def _extract_from_docx(self, docx_path):
        """Zamonaviy Word (.docx) fayldan matn o'qish"""
        try:
            doc = Document(docx_path)
            extracted_text = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            self.text = "\n".join(extracted_text)
            
            # Agar docx bo'sh bo'lsa yoki ichida rasm bo'lsa, OCR ham qo'shib yuborish mumkin
            if not self.text.strip():
                self._extract_from_doc_via_pdf(docx_path)
        except Exception:
            self._extract_from_doc_via_pdf(docx_path)

    def _extract_from_doc_via_pdf(self, file_path):
        """Eski .doc yoki boshqa formatlarni PyMuPDF qo'llab-quvvatlasa rasmga o'girib o'qiydi"""
        try:
            extracted_text = []
            doc = fitz.open(file_path)
            for page_index in range(len(doc)):
                page = doc[page_index]
                pix = page.get_pixmap(dpi=150)
                image = Image.open(BytesIO(pix.tobytes("png")))
                extracted_text.append(pytesseract.image_to_string(image, lang="eng"))
            self.text = "\n".join(extracted_text)
            doc.close()
        except Exception as e:
            self.text = f"Faylni o'qishda xatolik yuz berdi: {str(e)}"

    def _extract_from_text_or_csv(self, file_path):
        """Oddiy matn yoki CSV fayllar uchun"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            self.text = f.read()

    def _extract_from_excel(self, xlsx_path):
        """Excel (.xlsx) fayldagi barcha jadvallar matnini yig'ib olish"""
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        extracted_text = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                row_str = " ".join([str(cell) for cell in row if cell is not None])
                if row_str.strip():
                    extracted_text.append(row_str)
        self.text = "\n".join(extracted_text)